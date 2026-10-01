"""Normalize creators' free-text enrichment values to the fixed Lists vocabulary.

Deterministic: a fixed lookup table, no guessing. Any raw value missing from
its mapping table becomes UNMAPPED and is logged.
"""
import logging
import os
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv()

DATA_DIR = Path(__file__).parent / "data"
CREATORS_XLSX = os.getenv("CREATORS_XLSX", str(DATA_DIR / "creator_workbook.xlsx"))
MAPPINGS_XLSX = os.getenv("MAPPINGS_XLSX", str(DATA_DIR / "mapping_tables.xlsx"))

UNMAPPED = "UNMAPPED"

# raw column header -> (clean column name, Lists vocab header, mapping tab or None for 1:1 pass-through)
COLUMNS = {
    "Content Category": ("content_category", "ContentCategory", "AB_Category"),
    "Content Tone": ("content_tone", "ContentTone", "AC_Tone"),
    "Production Style": ("production_style", "ProductionStyle", "AD_Production"),
    "Regularly Features Pune Venues?": ("features_pune", "FeaturesFrequency", "AF_FeaturesPune"),
    "Prior F&B Brand Collabs?": ("prior_fnb_collabs", "YesNo", "AG_PriorCollabs"),
    "Engagement Authenticity Read": ("authenticity", "AuthenticityRead", None),
}

log = logging.getLogger(__name__)


def _clean(value):
    return value.strip() if isinstance(value, str) else value


def load_vocab(creators_xlsx=CREATORS_XLSX):
    """{Lists header: set of valid values} from the 'Lists' sheet."""
    ws = load_workbook(creators_xlsx, read_only=True, data_only=True)["Lists"]
    rows = list(ws.iter_rows(values_only=True))
    vocab = {
        header: {_clean(r[i]) for r in rows[1:] if r[i] is not None}
        for i, header in enumerate(rows[0]) if header
    }
    # Not in the Lists tab: features-city is a 3-level signal (partial proxy credit for "Occasionally").
    vocab["FeaturesFrequency"] = vocab["YesNo"] | {"Occasionally"}
    return vocab


def load_mappings(vocab, mappings_xlsx=MAPPINGS_XLSX):
    """{raw column header: {raw value: Lists value}}. Fails fast on an invalid target."""
    wb = load_workbook(mappings_xlsx, read_only=True, data_only=True)
    mappings = {}
    for column, (_, vocab_key, tab) in COLUMNS.items():
        valid = vocab[vocab_key]
        if tab is None:
            mappings[column] = {v: v for v in valid}
            continue
        rows = list(wb[tab].iter_rows(values_only=True))
        start = next(i for i, r in enumerate(rows) if str(r[0]).startswith("Raw value")) + 1
        table = {}
        for raw, target, *_ in rows[start:]:
            if raw is None:
                continue
            raw, target = _clean(raw), _clean(target)
            if target not in valid:
                raise ValueError(f"{tab}: {raw!r} maps to {target!r}, not in Lists/{vocab_key}")
            if table.get(raw, target) != target:
                raise ValueError(f"{tab}: {raw!r} mapped twice with different targets")
            table[raw] = target
        mappings[column] = table
    return mappings


def load_creators(creators_xlsx=CREATORS_XLSX):
    """Rows of 'Influencer Master' as dicts keyed by header; blank rows skipped."""
    ws = load_workbook(creators_xlsx, read_only=True, data_only=True)["Influencer Master"]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    return [dict(zip(headers, r)) for r in rows if any(v is not None for v in r)]


def normalize(creators, mappings):
    """Return copies of creators with clean columns added. Unknown raw values -> UNMAPPED (logged)."""
    out = []
    for creator in creators:
        row = dict(creator)
        for column, (clean_name, _, _) in COLUMNS.items():
            raw = _clean(creator.get(column))
            mapped = mappings[column].get(raw, UNMAPPED)
            if mapped == UNMAPPED:
                log.warning("UNMAPPED %s=%r for creator %s", column, raw, creator.get("Username"))
            row[clean_name] = mapped
        out.append(row)
    return out


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    vocab = load_vocab()
    rows = normalize(load_creators(), load_mappings(vocab))
    unmapped = sum(r[c] == UNMAPPED for r in rows for c, *_ in COLUMNS.values())
    print(f"Normalized {len(rows)} creators; {unmapped} UNMAPPED values.")
